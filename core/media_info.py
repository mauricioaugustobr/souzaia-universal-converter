import os
import json
import subprocess
from pathlib import Path
from typing import Dict, Any, List, Optional
from PIL import Image

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

# Extension to category mapping
EXTENSION_CATEGORIES = {
    # Video
    'mp4': 'video', 'mkv': 'video', 'avi': 'video', 'mov': 'video',
    'webm': 'video', 'flv': 'video', 'wmv': 'video', 'm4v': 'video',
    'mpg': 'video', 'mpeg': 'video', '3gp': 'video', 'ogv': 'video',
    'ts': 'video', 'vob': 'video',

    # Audio
    'mp3': 'audio', 'wav': 'audio', 'flac': 'audio', 'aac': 'audio',
    'ogg': 'audio', 'm4a': 'audio', 'opus': 'audio', 'wma': 'audio',
    'aiff': 'audio', 'alac': 'audio', 'amr': 'audio', 'ac3': 'audio',
    'dts': 'audio',

    # Image
    'png': 'image', 'jpg': 'image', 'jpeg': 'image', 'webp': 'image',
    'gif': 'image', 'bmp': 'image', 'tiff': 'image', 'tif': 'image',
    'ico': 'image', 'svg': 'image', 'tga': 'image', 'ppm': 'image',
    'eps': 'image', 'avif': 'image',

    # Document
    'pdf': 'document', 'docx': 'document', 'txt': 'document',
    'md': 'document', 'html': 'document', 'htm': 'document',
    'csv': 'document', 'xlsx': 'document', 'xls': 'document',
    'json': 'document', 'xml': 'document', 'rtf': 'document',
}

# Supported target formats by category
TARGET_FORMATS = {
    'video': {
        'video': ['mp4', 'mkv', 'webm', 'mov', 'avi', 'flv', 'wmv', 'gif'],
        'audio': ['mp3', 'wav', 'flac', 'aac', 'ogg', 'opus', 'm4a']
    },
    'audio': {
        'audio': ['mp3', 'wav', 'flac', 'aac', 'ogg', 'opus', 'm4a', 'aiff', 'wma']
    },
    'image': {
        'image': ['png', 'jpg', 'jpeg', 'webp', 'gif', 'bmp', 'tiff', 'ico', 'svg'],
        'document': ['pdf']
    },
    'document': {
        'document': ['pdf', 'docx', 'txt', 'md', 'html', 'csv', 'xlsx', 'json', 'xml'],
        'image': ['png', 'jpg', 'webp']
    }
}

def get_file_category(filename_or_ext: str) -> str:
    ext = filename_or_ext.lower().lstrip('.')
    return EXTENSION_CATEGORIES.get(ext, 'unknown')

def format_size(bytes_num: int) -> str:
    for unit in ['B', 'KB', 'MB', 'GB', 'TB']:
        if bytes_num < 1024.0:
            return f"{bytes_num:.1f} {unit}" if unit != 'B' else f"{bytes_num} B"
        bytes_num /= 1024.0
    return f"{bytes_num:.1f} PB"

def inspect_media(filepath: str) -> Dict[str, Any]:
    path = Path(filepath)
    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {filepath}")

    ext = path.suffix.lower().lstrip('.')
    category = get_file_category(ext)
    size_bytes = path.stat().st_size
    info: Dict[str, Any] = {
        'filename': path.name,
        'extension': ext,
        'category': category,
        'size_bytes': size_bytes,
        'size_formatted': format_size(size_bytes),
        'details': {},
        'target_formats': []
    }

    # Available targets based on category
    if category == 'video':
        info['target_formats'] = TARGET_FORMATS['video']['video'] + TARGET_FORMATS['video']['audio']
        _inspect_av_file(filepath, info)
    elif category == 'audio':
        info['target_formats'] = TARGET_FORMATS['audio']['audio']
        _inspect_av_file(filepath, info)
    elif category == 'image':
        info['target_formats'] = TARGET_FORMATS['image']['image'] + TARGET_FORMATS['image']['document']
        _inspect_image_file(filepath, info)
    elif category == 'document':
        info['target_formats'] = TARGET_FORMATS['document']['document'] + TARGET_FORMATS['document']['image']
        _inspect_document_file(filepath, info)
    else:
        # Fallback to general formats
        info['target_formats'] = ['pdf', 'txt', 'mp4', 'mp3', 'png', 'jpg']

    return info

def _inspect_av_file(filepath: str, info: Dict[str, Any]):
    try:
        cmd = [
            'ffprobe', '-v', 'quiet',
            '-print_format', 'json',
            '-show_format', '-show_streams',
            filepath
        ]
        result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, check=True)
        probe_data = json.loads(result.stdout)
        
        streams = probe_data.get('streams', [])
        format_info = probe_data.get('format', {})

        video_stream = next((s for s in streams if s.get('codec_type') == 'video'), None)
        audio_stream = next((s for s in streams if s.get('codec_type') == 'audio'), None)

        details: Dict[str, Any] = {
            'duration_sec': float(format_info.get('duration', 0.0)),
            'bitrate_kbps': int(format_info.get('bit_rate', 0)) // 1000 if format_info.get('bit_rate') else None,
        }

        if details['duration_sec']:
            mins = int(details['duration_sec'] // 60)
            secs = int(details['duration_sec'] % 60)
            details['duration_formatted'] = f"{mins:02d}:{secs:02d}"

        if video_stream:
            details['video_codec'] = video_stream.get('codec_name')
            details['width'] = video_stream.get('width')
            details['height'] = video_stream.get('height')
            details['resolution'] = f"{details['width']}x{details['height']}"
            r_frame_rate = video_stream.get('r_frame_rate', '0/1')
            if '/' in r_frame_rate:
                num, den = map(float, r_frame_rate.split('/'))
                fps = round(num / den, 2) if den > 0 else 0
                details['fps'] = fps

        if audio_stream:
            details['audio_codec'] = audio_stream.get('codec_name')
            details['channels'] = audio_stream.get('channels')
            details['sample_rate_hz'] = audio_stream.get('sample_rate')

        info['details'] = details
    except Exception as e:
        info['details'] = {'error_inspecting': str(e)}

def _inspect_image_file(filepath: str, info: Dict[str, Any]):
    try:
        with Image.open(filepath) as img:
            details = {
                'width': img.width,
                'height': img.height,
                'resolution': f"{img.width}x{img.height}",
                'format': img.format,
                'mode': img.mode,
                'animated': getattr(img, "is_animated", False),
                'n_frames': getattr(img, "n_frames", 1)
            }
            if 'dpi' in img.info:
                dpi = img.info['dpi']
                details['dpi'] = f"{int(dpi[0])}x{int(dpi[1])}" if isinstance(dpi, tuple) else str(dpi)
            info['details'] = details
    except Exception as e:
        info['details'] = {'error_inspecting': str(e)}

def _inspect_document_file(filepath: str, info: Dict[str, Any]):
    ext = info['extension']
    try:
        if ext == 'pdf':
            doc = pymupdf.open(filepath)
            details = {
                'pages': len(doc),
                'title': doc.metadata.get('title') or '',
                'author': doc.metadata.get('author') or '',
                'encrypted': doc.is_encrypted
            }
            doc.close()
            info['details'] = details
        elif ext in ['csv', 'txt', 'md', 'html', 'json', 'xml']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                lines = sum(1 for _ in f)
            info['details'] = {'line_count': lines}
        elif ext in ['xlsx', 'xls']:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, read_only=True)
                info['details'] = {'sheets': wb.sheetnames, 'total_sheets': len(wb.sheetnames)}
                wb.close()
            except Exception:
                pass
    except Exception as e:
        info['details'] = {'error_inspecting': str(e)}
